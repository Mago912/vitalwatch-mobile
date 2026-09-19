from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "VitalWatch_Registro_Implementacion_Mensajeria.docx"

NAVY = "17365D"
PALE_BLUE = "EEF5FB"
PALE_GRAY = "F7F7F7"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"
RED = "C00000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LIGHT_BORDER)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def set_repeat_header_and_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.text = "VITALWATCH   |   REGISTRO TÉCNICO DE IMPLEMENTACIÓN"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.bold = True
    add_page_number(section.footer.paragraphs[0])


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    set_repeat_header_and_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor(0, 0, 0)

    for name, size, before, after in (
        ("Heading 1", 18, 12, 8),
        ("Heading 2", 13, 10, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(10.2)
        styles[name].paragraph_format.space_after = Pt(3)


def add_paragraph(doc, text="", bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbers(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.32)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{index}.  ").bold = True
        p.add_run(item)


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_table(doc, headers, rows, widths=None, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False if widths else True
    table.style = "Table Grid"
    set_table_borders(table)
    header = table.rows[0]
    repeat_table_header(header)
    for index, title in enumerate(headers):
        cell = header.cells[index]
        cell.text = title
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[index])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(font_size)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[index])
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, WHITE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor(0, 0, 0)
    return table


def add_change_record(doc, change_id, file_name, existed, added, minimum, reason, preserved):
    doc.add_heading(f"{change_id}  {file_name}", level=2)
    add_table(
        doc,
        ["Campo", "Registro"],
        [
            ("Qué existía", existed),
            ("Qué se agregó", added),
            ("Cambio mínimo", minimum),
            ("Por qué", reason),
            ("Qué se preservó", preserved),
        ],
        widths=[1.35, 5.75],
        font_size=8.8,
    )


def part(doc, number, title):
    heading = doc.add_heading(f"PARTE {number}  {title}", level=1)
    heading.paragraph_format.page_break_before = True


def build_document():
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Registro de implementación de Mensajería en VitalWatch")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("App 1.0.9 y BIOSYS 1.0.7")
    subtitle2 = doc.add_paragraph()
    subtitle2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle2.add_run("CANDIDATE FOR REVIEW")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(192, 0, 0)

    doc.add_paragraph()
    add_table(
        doc,
        ["Dato", "Valor"],
        [
            ("Proyecto", "VitalWatch existente"),
            ("Fecha de cierre técnico", "14 de septiembre de 2026"),
            ("Versión de app", "1.0.9 candidata"),
            ("Firmware integrado", "BIOSYS 1.0.7 = SYS 0.9.7 + BIO 0.6.3"),
            ("Última baseline física conocida", "BIOSYS 1.0.5"),
            ("Backend", "Supabase Auth, PostgreSQL con RLS, Realtime y Edge Functions"),
            ("Estado", "Implementado y validado estáticamente; falta prueba integral en equipos reales"),
        ],
        widths=[2.05, 5.05],
        font_size=9.2,
    )
    add_paragraph(
        doc,
        "Propósito. Este registro permite que el equipo estudiantil identifique qué se cambió, por qué se cambió, qué quedó intacto, cómo probar el candidato y cómo volver al estado anterior sin confundir una compilación correcta con una validación clínica o física.",
        "Propósito.",
    )
    add_paragraph(
        doc,
        "Conclusión principal. El delta añade Mensajería como primera opción de la pulsera, contactos administrados por la app y eventos auditables en Supabase. Los números telefónicos no se descargan al ESP32, las notificaciones de pantalla bloqueada son genéricas y la llamada o el mensaje sólo continúan mediante las aplicaciones del sistema operativo. No se modificaron algoritmos biomédicos ni GPIO.",
        "Conclusión principal.",
    )

    doc.add_page_break()
    doc.add_heading("Contenido y lectura del estado", level=1)
    add_numbers(doc, [
        "Estado inicial",
        "Firmware y pulsera",
        "Aplicación",
        "Supabase",
        "Notificaciones",
        "Telegram",
        "Seguridad y privacidad",
        "Integración completa",
        "Archivos modificados",
        "Archivos preservados",
        "Pruebas",
        "Cómo probar nosotros",
        "Problemas encontrados",
        "Pendientes",
        "Rollback",
    ])
    add_paragraph(doc, "Cómo leer los resultados. PASS significa que se ejecutó la comprobación indicada. NOT TESTED significa que la implementación existe, pero esa conducta necesita teléfono, placa, cuenta adicional o intervención humana. No se utiliza un PASS de Android como sustituto de iOS, ni una compilación como sustituto de una prueba física.")

    part(doc, 1, "Estado inicial")
    add_paragraph(doc, "La inspección confirmó que VitalWatch no es un proyecto nuevo. Ya reunía firmware Arduino para ESP32, una app Expo y React Native, Supabase como backend, autenticación, telemetría, control TFT, medicación, historial, push Expo y Telegram opcional. La implementación se hizo de forma aditiva sobre esas piezas.")
    add_table(doc, ["Área", "Estado encontrado"], [
        ("Versión más reciente previa", "App 1.0.8 y BIOSYS 1.0.6 = SYS 0.9.6 + BIO 0.6.3"),
        ("Baseline física", "BIOSYS 1.0.5; se mantiene como última referencia instalada/probada conocida"),
        ("Firmware", "Arduino para ESP32 Dev Module; fuentes .ino, .cpp y .h; TFT ST7735; WiFi y HTTPS"),
        ("App", "Expo SDK 54, React Native 0.81.5, Expo Router 6 y TypeScript"),
        ("Android", "Paquete com.vitalwatchmobile; notificaciones Expo; Google services configurado"),
        ("iOS", "Bundle com.vitalwatchmobile; rutas y Linking compartidos; soporte de tablet activo"),
        ("Supabase", "Auth, PostgreSQL, RLS, Realtime, Edge Functions y tokens push"),
        ("Comunicación real", "ESP32 a Supabase por funciones HTTPS; app a Supabase por SDK; no se inventó Bluetooth"),
        ("Telegram", "Bot, enlace temporal y contactos Telegram existentes; canal opcional"),
        ("Riesgo de repositorio", "Git contiene objetos faltantes y cambios previos no relacionados; no es una fuente de rollback confiable"),
    ], widths=[1.65, 5.45], font_size=8.9)
    doc.add_heading("Decisión de versionado", level=2)
    add_bullets(doc, [
        "Working target: App 1.0.9 y BIOSYS 1.0.7.",
        "Composición: SYS 0.9.7 para el sistema y BIO 0.6.3 preservado.",
        "La candidata previa BIOSYS 1.0.6 ya compilaba, pero no estaba probada físicamente.",
        "No se promovió 1.0.6 ni 1.0.7 a baseline. La etiqueta final sigue siendo CANDIDATE FOR REVIEW.",
    ])
    doc.add_heading("Funciones reutilizadas", level=2)
    add_bullets(doc, [
        "Tabla emergency_contacts, ampliada previamente para Telegram y teléfono.",
        "Tabla device_events e índice único device_id + source_event_id.",
        "Supabase Auth y la relación de propiedad entre usuario y dispositivo.",
        "Proveedor de estado de la app y sus suscripciones Realtime.",
        "Tarea de red del firmware, navegación con tres botones y estilo TFT existente.",
        "Registro Expo Push y Edge Function send-vitalwatch-push.",
        "Bot y enlace temporal de Telegram.",
    ])

    part(doc, 2, "Firmware y pulsera")
    add_paragraph(doc, "La carpeta BIOSYS 1.0.7 se creó a partir de BIOSYS 1.0.6. Sólo se editaron los puntos necesarios para versiones, menú, interfaz y red. Los archivos biomédicos y de botones se copiaron sin diferencias binarias verificables.")
    add_change_record(doc, "VW-MSG-01", "Configuracion.h e Interfaz.h", "El menú principal tenía cuatro opciones: Signos, Movimiento, Estado y Medicación.", "MENSAJERÍA como primera opción y un nuevo modo visual.", "Se aumentó el número de opciones de cuatro a cinco y se desplazaron los índices existentes.", "Cumplir el acceso directo solicitado sin crear otra navegación.", "Orden relativo de las cuatro opciones anteriores, colores, botones, GPIO y flujo de alertas.")
    add_change_record(doc, "VW-MSG-02", "Mensajeria.h", "No había sincronización de agenda ni solicitudes Mensaje/Llamada desde la pulsera.", "Módulo acotado a ocho contactos con ID y nombre; cola RAM de cuatro solicitudes; descarga y envío HTTPS.", "Se añadió un encabezado independiente y se reutilizó la configuración de red y credencial del dispositivo.", "Evitar que números telefónicos o secretos de usuario residan en la pulsera.", "Bucle biomédico, sensores, tareas existentes y bibliotecas.")
    add_change_record(doc, "VW-MSG-03", "Sincronizacion_Medicacion.h", "Una tarea de red existente sincronizaba medicación, controles TFT y telemetría.", "Inicialización y procesamiento periódico no bloqueante de Mensajeria.", "Dos puntos de integración en la tarea existente; no se creó una tarea adicional.", "Evitar bloquear loop y reducir competencia por CPU/red con sensores.", "Frecuencias de muestreo, acceso I2C, envío de telemetría y lógica de medicamentos.")
    add_change_record(doc, "VW-MSG-04", "Interfaz.h y VitalWatch_BIOSYS_1_0_7.ino", "La TFT ya usaba izquierda, derecha, OK y pulsación larga para navegar.", "Flujo Tipo, Contacto, Confirmación y Resultado; estados Sin WiFi, Actualizando, Pendiente, Aceptado y Error.", "Se añadieron pantallas y dispatch del nuevo modo sin reemplazar las funciones anteriores.", "La pulsera solicita una acción, pero no puede verificar que el teléfono haya enviado o llamado.", "Prioridad de alertas, cierre físico de impacto, resto del menú y comportamiento de botones.")
    doc.add_heading("Flujo visible", level=2)
    add_code_block(doc, [
        "MENSAJERÍA",
        "  -> MENSAJE | LLAMADA",
        "  -> CONTACTO (nombre, nunca teléfono)",
        "  -> CONFIRMAR",
        "  -> PENDIENTE | ACEPTADO | ERROR DE RED",
        "  -> 'No implica envío'",
    ])
    add_paragraph(doc, "La palabra ACEPTADO significa que Supabase aceptó/encoló el evento. No significa que el SMS se envió ni que la llamada comenzó. Si el ESP32 pierde red, reintenta las solicitudes conservadas en RAM. Un reinicio puede perder una solicitud no confirmada; este límite queda pendiente para una versión posterior.")

    part(doc, 3, "Aplicación")
    doc.add_heading("Código común", level=2)
    add_bullets(doc, [
        "constants/vitalwatch.ts declara las nuevas versiones y metadatos de eventos/contactos.",
        "lib/vitalwatch-emergency-contacts.ts implementa alta, edición, borrado y activación bajo la sesión actual.",
        "providers/vitalwatch-provider.tsx integra contactos, estado de sincronización, historial y Realtime.",
        "lib/vitalwatch-api.ts consulta el detalle del evento y recién entonces resuelve el contacto bajo RLS.",
    ])
    doc.add_heading("Contactos y sincronización", level=2)
    add_paragraph(doc, "La nueva pantalla Contactos administra una agenda propia de VitalWatch. No solicita permiso para leer toda la agenda del teléfono. Cada cambio se guarda en Supabase; la pulsera obtiene en su próxima sincronización sólo los contactos activos, telefónicos y con máximo ocho registros.")
    add_table(doc, ["Acción", "App", "Supabase", "ESP32"], [
        ("Añadir", "Valida nombre y teléfono", "Inserta bajo user_id de sesión", "Aparece como ID + nombre"),
        ("Editar", "Actualiza campos", "RLS limita el registro", "Recibe nombre actualizado"),
        ("Desactivar", "Conserva registro", "active = false", "Deja de aparecer"),
        ("Eliminar", "Solicita acción explícita", "Borra bajo RLS", "Deja de aparecer; eventos conservan contexto disponible"),
    ], widths=[1.05, 2.0, 2.05, 2.0], font_size=8.3)
    doc.add_heading("Historial y detalle", level=2)
    add_paragraph(doc, "Historial continúa mostrando los eventos almacenados. message_request aparece como MENSAJE y call_request como LLAMADA, con el nombre autorizado cuando está disponible. fall_detected se presenta como CAÍDA en texto y rojo; por lo tanto no depende sólo del color.")
    doc.add_heading("Navegación", level=2)
    add_paragraph(doc, "Las rutas contacts y event/[eventId] son pantallas internas dentro del grupo protegido. No se añadieron botones a la barra inferior. Una notificación abre /event/<id>; si faltan sesión o vínculo, la ruta se guarda en AsyncStorage y se consume cuando el acceso queda listo.")
    doc.add_heading("Llamadas y mensajes", level=2)
    add_paragraph(doc, "El detalle autorizado recupera el teléfono sólo dentro de la app. Para LLAMADA se comprueba tel: y se abre el marcador; para MENSAJE se comprueba sms: y se abre el compositor. Android y iOS conservan sus confirmaciones y restricciones nativas. VitalWatch no realiza el envío automáticamente.")
    doc.add_heading("Android", level=2)
    add_bullets(doc, [
        "Se reutiliza expo-notifications y el canal vitalwatch-alerts.",
        "Se exportó el bundle Android con Expo SDK 54 y se generó un APK interno firmado 1.0.9, versionCode 17, mediante EAS.",
        "Desde Expo SDK 53, push remota Android requiere development build o binario nativo; Expo Go no basta.",
        "El APK quedó verificado localmente, pero la apertura real de tel:, sms:, lock screen y toque requiere un teléfono Android conectado.",
    ])
    doc.add_heading("iOS", level=2)
    ios_items = [
        "Se usa el mismo intent de usuario mediante expo-linking y el listener de expo-notifications.",
        "Se exportó el bundle iOS por separado; no se infirió PASS desde Android.",
        "APNs, permisos, lock screen, toque con app cerrada y apertura tel:/sms: requieren un build firmado y dispositivo iOS.",
    ]
    for index, item in enumerate(ios_items):
        paragraph = doc.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.keep_with_next = index < len(ios_items) - 1

    part(doc, 4, "Supabase")
    add_paragraph(doc, "Supabase continúa como única arquitectura backend. No se añadió Firebase como base de datos. Expo, FCM o APNs sólo actúan como transporte de notificaciones donde corresponde.")
    add_table(doc, ["Componente", "Decisión"], [
        ("Auth", "Se preservó Supabase Auth; no se creó otro login."),
        ("emergency_contacts", "Reutilizada para teléfono y Telegram; no se duplicó una tabla de agenda."),
        ("device_events", "Se agregó contact_id opcional con FK ON DELETE SET NULL."),
        ("RLS", "La app sigue consultando con la sesión del usuario; las políticas existentes filtran por propietario."),
        ("Realtime", "Se preservó para actualizar historial/estado con la app abierta."),
        ("Edge Functions", "Nueva vitalwatch-device-messaging; send-vitalwatch-push ajustada; Telegram preservado."),
    ], widths=[1.65, 5.45], font_size=9)
    doc.add_heading("Migración VW-SUPA-01 y VW-SEC-01", level=2)
    add_paragraph(doc, "La migración 20260914190000_messaging_contacts_and_events.sql añade la relación y un índice parcial. Un trigger SECURITY DEFINER valida que el contacto esté activo y pertenezca al mismo usuario que el dispositivo. La función fija search_path, revoca ejecución pública y la concede a service_role.")
    doc.add_heading("Edge Function VW-SUPA-02 y VW-SEC-02", level=2)
    add_bullets(doc, [
        "POST action=list autentica deviceCode + x-device-token y retorna como máximo ocho objetos id/displayName.",
        "phone_e164 se usa sólo como filtro de capacidad y nunca forma parte de la respuesta al ESP32.",
        "POST action=request valida evento, fecha, tipo, contacto activo, canal telefónico y propietario.",
        "El evento se guarda en minúsculas como message_request o call_request, conservando el contrato de la base.",
        "El upsert usa device_id + source_event_id con ignoreDuplicates para aceptar reintentos sin duplicar registros.",
        "Las respuestas establecen Cache-Control no-store y los errores externos son genéricos.",
    ])
    add_paragraph(doc, "Estado remoto. La migración figura con el mismo identificador local y remoto. vitalwatch-device-messaging y send-vitalwatch-push fueron desplegadas desde el panel de Supabase. La verificación JWT heredada quedó desactivada sólo en la función del ESP32 porque su autenticación real es la credencial de dispositivo validada dentro de la función.")

    part(doc, 5, "Notificaciones")
    doc.add_heading("Recorrido", level=2)
    add_code_block(doc, [
        "EVENTO EN device_events",
        "  -> send-vitalwatch-push",
        "  -> Expo Push (FCM/APNs como transporte)",
        "  -> lock screen con texto genérico",
        "  -> toque del usuario",
        "  -> control de sesión/vinculación",
        "  -> /event/<event_id>",
        "  -> consulta RLS y detalle autorizado",
    ])
    add_change_record(doc, "VW-NOTIF-01 y VW-NOTIF-04", "send-vitalwatch-push e implementación local", "El push podía incluir el texto detallado del evento en la pantalla bloqueada.", "Título VitalWatch, textos genéricos y data mínima event_id, event_type, url.", "Se cambió sólo la composición del mensaje Expo; el registro de tokens y transporte permanecen.", "Evitar exposición de signos, nombres, números o historia médica sin desbloquear.", "Expo Push, tokens por usuario, prioridad, sonido y Telegram de emergencias.")
    add_change_record(doc, "VW-NOTIF-02 y VW-NOTIF-03", "vitalwatch-notification-routing.ts y app/_layout.tsx", "El toque no tenía un destino robusto al evento exacto a través de login/vinculación.", "Lista blanca /event/<id>, listener, lectura de la última respuesta y ruta pendiente.", "Se añadieron dos efectos al navegador raíz protegido.", "Impedir rutas arbitrarias desde un payload manipulado y conservar la intención del usuario.", "Navegación existente, login, vinculación y barra de pestañas.")
    add_paragraph(doc, "Los detalles de contacto se consultan en la pantalla de evento. Ni el teléfono ni el nombre viajan en la data del push. Un ID por sí solo no concede acceso: la consulta posterior depende de la sesión y RLS.")

    part(doc, 6, "Telegram")
    add_paragraph(doc, "Telegram ya existía con bot, contactos del canal telegram, token temporal y enlace profundo. Se preservó como canal opcional y separado de los contactos telefónicos.")
    add_bullets(doc, [
        "El botón de conexión abre t.me/<bot>?start=<código temporal> y conserva /vincular como alternativa manual.",
        "La pantalla Configuración filtra los contactos Telegram; la nueva pantalla Contactos muestra sólo la agenda telefónica.",
        "El push genérico no reemplaza Telegram. Las alertas de emergencia configuradas continúan por ese canal.",
        "MESSAGE_REQUEST y CALL_REQUEST no se clasifican como emergencia médica ni se fuerzan por Telegram.",
        "No se cambió ni expuso el token del bot en app o firmware.",
    ])
    add_paragraph(doc, "Mejora mínima de esta entrega. La separación visual y de modelo evita que un contacto Telegram sin teléfono sea descargado por la pulsera, y evita que borrar un contacto telefónico altere accidentalmente el flujo del bot.")

    part(doc, 7, "Seguridad y privacidad")
    add_table(doc, ["Riesgo", "Mitigación implementada", "Estado"], [
        ("Teléfono en la pulsera", "Respuesta de contactos limitada a id + displayName", "PASS remoto"),
        ("ID de contacto manipulado", "Validación Edge por usuario/canal/active y trigger de base", "PASS remoto"),
        ("Push sensible en lock screen", "Texto genérico y payload mínimo", "Verificado en código; hardware pendiente"),
        ("Ruta arbitraria desde push", "Regex estricta /event/<entero positivo>", "Verificado estáticamente"),
        ("Acceso sin sesión", "Rutas protegidas y RLS", "401 remoto en endpoints probados"),
        ("Secreto administrativo en cliente", "No se agregó service-role ni token Telegram a app/firmware", "Búsqueda estática PASS"),
        ("Duplicado por reintento", "Índice único y upsert ignoreDuplicates", "Diseño verificado; reintento real pendiente"),
        ("Permisos excesivos", "Agenda propia; no se solicita acceso a contactos del teléfono", "PASS por inspección"),
        ("Afirmación falsa de envío", "Estados Pendiente/Aceptado/Error y texto explicativo", "PASS por inspección"),
    ], widths=[1.55, 4.35, 1.2], font_size=8.2)
    doc.add_heading("Datos sensibles y límites", level=2)
    add_paragraph(doc, "Se consideran sensibles el número telefónico, los signos vitales, datos biomédicos, nombres y credenciales. El número permanece en Supabase y se obtiene dentro de la app autenticada sólo al abrir el evento. El firmware conserva una credencial de dispositivo en su archivo privado local, que se excluyó del ZIP Arduino.")
    doc.add_heading("Autorización", level=2)
    add_paragraph(doc, "La autorización se comprueba en dos capas. La Edge Function enlaza el dispositivo autenticado con su user_id y vuelve a consultar el contacto. La base repite la comprobación antes de insertar o cambiar device_id/contact_id. Esto evita confiar únicamente en un ID del cliente.")
    doc.add_heading("Pruebas aún necesarias", level=2)
    add_paragraph(doc, "No se creó una segunda cuenta real durante esta entrega. La prueba remota rechazó un contact_id manipulado, y el trigger implementa la defensa cruzada, pero el escenario Usuario A frente a un contacto real de Usuario B debe ensayarse antes de promoción.")

    part(doc, 8, "Integración completa")
    add_code_block(doc, [
        "PULSERA ESP32",
        "  |  HTTPS: ID + nombre / solicitud mínima",
        "  v",
        "SUPABASE EDGE FUNCTION",
        "  |  autentica dispositivo y autoriza contacto",
        "  v",
        "POSTGRESQL + RLS + REALTIME",
        "  |  evento y notificación genérica",
        "  v",
        "APP ANDROID / IOS",
        "  |  sesión válida, detalle y teléfono autorizado",
        "  +--> marcador tel: o compositor sms:",
        "  +--> historial con MENSAJE, LLAMADA o CAÍDA",
        "  +--> Telegram opcional para alertas existentes",
    ])
    doc.add_heading("MESSAGE_REQUEST", level=2)
    add_paragraph(doc, "La persona selecciona Mensaje y un nombre en la TFT. El ESP32 crea un source_event_id, lo coloca en la cola y envía contactId y eventType. Supabase valida, registra message_request y produce una push genérica. Al tocarla, la app consulta el evento; el botón Abrir mensaje inicia el compositor con el número autorizado. El envío depende del usuario y del sistema operativo.")
    doc.add_heading("CALL_REQUEST", level=2)
    add_paragraph(doc, "El recorrido es equivalente, pero registra call_request. La app ofrece Abrir marcador después de comprobar que el contacto continúa activo y que la plataforma admite tel:. La pulsera nunca llama directamente.")
    doc.add_heading("FALL y SOS", level=2)
    add_paragraph(doc, "Los módulos existentes continúan generando sus eventos. No se cambiaron umbrales, impacto, inmovilidad ni muestreo. FALL llega a Supabase, recibe push genérica y se ve en Historial como CAÍDA en rojo y texto. SOS conserva su flujo anterior.")
    doc.add_heading("Sincronización y errores", level=2)
    add_paragraph(doc, "El alta, edición, desactivación o borrado se refleja en Supabase. El siguiente ciclo de la pulsera descarga la lista activa. Sin WiFi, la TFT muestra el estado correspondiente; una solicitud aceptada por backend no se confunde con una acción completada en el teléfono.")

    part(doc, 9, "Archivos modificados")
    add_paragraph(doc, "La siguiente tabla identifica el delta de Mensajería. Archivos que ya tenían cambios anteriores del proyecto se tocaron sólo en los puntos indicados.")
    modified_rows = [
        ("App", "app.json; package.json; package-lock.json", "Modificar", "Versión candidata 1.0.9"),
        ("App", "constants/vitalwatch.ts", "Modificar", "Versiones, contactos y metadatos de eventos"),
        ("App", "lib/vitalwatch-emergency-contacts.ts", "Agregar", "CRUD y activación de agenda VitalWatch"),
        ("App", "providers/vitalwatch-provider.tsx", "Modificar", "Estado, sync, historial y Realtime"),
        ("App", "app/(tabs)/contacts.tsx", "Agregar", "Pantalla de contactos"),
        ("App", "app/(tabs)/settings.tsx", "Modificar", "Acceso a Contactos y separación Telegram"),
        ("App", "app/(tabs)/explore.tsx", "Modificar", "Eventos presionables y CAÍDA visible"),
        ("App", "app/(tabs)/event/[eventId].tsx", "Agregar", "Detalle autorizado y tel:/sms:"),
        ("App", "app/(tabs)/_layout.tsx", "Modificar", "Rutas internas protegidas"),
        ("App", "lib/vitalwatch-api.ts", "Modificar", "contact_id y consulta de detalle"),
        ("Notificaciones", "lib/vitalwatch-notifications.ts", "Modificar", "Texto local genérico"),
        ("Notificaciones", "lib/vitalwatch-notification-routing.ts", "Agregar", "Ruta validada y destino pendiente"),
        ("Notificaciones", "app/_layout.tsx", "Modificar", "Tap al evento mediante auth/vinculación"),
        ("Supabase", "20260914190000_messaging_contacts_and_events.sql", "Agregar", "FK, índice y trigger de autorización"),
        ("Supabase", "vitalwatch-device-messaging/index.ts", "Agregar", "Lista mínima y solicitudes idempotentes"),
        ("Supabase", "vitalwatch-device-messaging/deno.json", "Agregar", "Configuración de función"),
        ("Supabase", "send-vitalwatch-push/index.ts", "Modificar", "Push genérica y data mínima"),
        ("Supabase", "supabase/config.toml", "Modificar", "Registro de la función de dispositivo"),
        ("Firmware", "VitalWatch_BIOSYS_1_0_7/Configuracion.h", "Modificar", "Versiones y modo MENSAJERÍA"),
        ("Firmware", "VitalWatch_BIOSYS_1_0_7/Mensajeria.h", "Agregar", "Contactos, cola y red"),
        ("Firmware", "VitalWatch_BIOSYS_1_0_7/Interfaz.h", "Modificar", "Menú y flujo TFT"),
        ("Firmware", "VitalWatch_BIOSYS_1_0_7/Sincronizacion_Medicacion.h", "Modificar", "Integración en tarea de red"),
        ("Firmware", "VitalWatch_BIOSYS_1_0_7.ino", "Modificar", "Dispatch y procesamiento UI"),
        ("Firmware", "README.md; CAMBIOS_1_0_7.md; MANIFEST.md", "Modificar/Agregar", "Uso, cambios y evidencia"),
        ("Herramientas", "scripts/esp32-firmware.ps1", "Modificar", "Target BIOSYS 1.0.7"),
        ("Pruebas", "scripts/test-vitalwatch-security.mjs", "Modificar", "Casos de Mensajería y minimización"),
        ("Proyecto", "README.md; CODEX_HANDOFF.md; tsconfig.json", "Modificar", "Estado, continuidad y exclusión del backup"),
        ("Documentación", "VITALWATCH_MENSAJERIA_1_0_9_CAMBIOS.md", "Agregar", "Resumen operativo"),
    ]
    add_table(doc, ["Área", "Archivo", "Acción", "Motivo"], modified_rows, widths=[1.1, 3.25, 1.05, 2.0], font_size=7.5)

    part(doc, 10, "Archivos preservados")
    add_paragraph(doc, "Se calcularon SHA-256 sobre BIOSYS 1.0.6 y BIOSYS 1.0.7. Los nueve archivos protegidos siguientes tienen el mismo hash en ambas carpetas.")
    hashes = [
        ("Sensor_Oxigeno.cpp", "0040B145...E18357", "Igual"),
        ("Sensor_Oxigeno.h", "99DBCF07...ADCCB", "Igual"),
        ("Sensor_Movimiento.cpp", "AE469EDE...2CE3EF", "Igual"),
        ("Sensor_Movimiento.h", "378BA423...4A051A", "Igual"),
        ("I2CBusService.cpp", "32875DA0...7225D58", "Igual"),
        ("I2CBusService.h", "F64544EC...301CFBA", "Igual"),
        ("BioResearch.cpp", "2E246A20...5B8840D", "Igual"),
        ("BioResearch.h", "2206E99D...04DA2EA", "Igual"),
        ("Botones.h", "A242CA23...81B7CA7", "Igual"),
    ]
    add_table(doc, ["Archivo", "SHA-256 abreviado", "Resultado"], hashes, widths=[2.65, 3.15, 1.3], font_size=8.8)
    add_bullets(doc, [
        "MAX30102, BPM, SpO2, PPG y filtros: sin cambios.",
        "MPU, acelerómetro, giroscopio, movimiento, impacto, caída e inmovilidad: sin cambios.",
        "GPIO y temporización biomédica: sin cambios.",
        "Botones físicos: archivo idéntico; sólo se consumen dentro del nuevo modo visual.",
        "BIOSYS 1.0.6 permanece completa como base inmediata; BIOSYS 1.0.5 permanece como referencia física.",
    ])
    add_paragraph(doc, "La compilación creció aproximadamente 7.512 bytes de flash respecto del tamaño documentado de la candidata 1.0.6 y 248 bytes de RAM global. El uso de flash queda en 89 %, por lo que futuras ampliaciones deben priorizar medición y recorte antes de añadir recursos grandes.")

    part(doc, 11, "Pruebas")
    tests = [
        ("Firmware ESP32", "PASS", "Compilación Arduino: 1.177.376 B flash (89 %), 54.704 B RAM (16 %)"),
        ("Binario", "PASS", "1.177.520 B; SHA-256 0E976BC8...FB4B65"),
        ("Sensores protegidos", "PASS", "9/9 hashes iguales entre BIOSYS 1.0.6 y 1.0.7"),
        ("TypeScript", "PASS", "npx tsc --noEmit; exit 0"),
        ("Lint", "PASS", "npm run lint; exit 0"),
        ("Lecturas", "PASS", "npm run test:readings; 7/7"),
        ("Telegram", "PASS", "npm run test:telegram; 5/5"),
        ("Seguridad remota", "PASS", "test:security: auth, token, payload mínimo, contacto manipulado"),
        ("Migración", "PASS", "20260914190000 presente local y remoto"),
        ("Edge messaging", "PASS", "Función desplegada; lista con token válido HTTP 200 e inválido HTTP 401"),
        ("Push backend", "PASS", "send-vitalwatch-push actualizado y desplegado"),
        ("Android bundle", "PASS", "Expo export Android: 1.513 módulos, bundle 4,58 MB"),
        ("APK Android", "PASS", "EAS build e257c85e...9789; 83.967.864 B; SHA-256 AA8DA345...7E2B"),
        ("iOS bundle", "PASS", "Expo export iOS: 1.498 módulos, bundle 4,51 MB"),
        ("ESP32 físico", "NOT TESTED", "No se cargó BIOSYS 1.0.7 ni se recorrió la TFT real"),
        ("Android real", "NOT TESTED", "APK 1.0.9 build 17 generado; sin dispositivo ADB para instalar/probar push/tap/tel/sms"),
        ("iOS real", "NOT TESTED", "Sin build firmado/dispositivo; push/tap/tel/sms pendientes"),
        ("Dos usuarios reales", "NOT TESTED", "Trigger implementado y ID manipulado rechazado; falta A contra B real"),
        ("Idempotencia real", "NOT TESTED", "Índice/upsert revisados; no se generó una alerta duplicada de prueba"),
        ("Validación clínica", "NOT TESTED", "Fuera del alcance; algoritmos no se modificaron"),
    ]
    add_table(doc, ["Área", "Estado", "Evidencia"], tests, widths=[1.65, 1.05, 4.4], font_size=8.1)
    add_paragraph(doc, "No hubo resultados FAIL en las comprobaciones ejecutadas. Los elementos NOT TESTED son condiciones de aceptación física y multiplataforma que no pueden declararse aprobadas con export o inspección de código.")

    part(doc, 12, "Cómo probar nosotros")
    doc.add_heading("Preparación", level=2)
    add_numbers(doc, [
        "Conservar BIOSYS 1.0.5 como referencia y guardar el ZIP de rollback.",
        "No compartir .env.local, vitalwatch_config.h, claves de Supabase ni token de Telegram.",
        "Instalar docs/VitalWatch_APP_1.0.9_CANDIDATE_build17.apk para la prueba Android; no usar Expo Go.",
        "Crear o elegir un contacto telefónico de prueba que haya autorizado el ensayo.",
    ])
    doc.add_heading("Prueba de la app", level=2)
    add_numbers(doc, [
        "Abrir App 1.0.9, iniciar sesión y verificar la pulsera vinculada.",
        "En Configuración, abrir Contactos. Añadir nombre y teléfono de prueba.",
        "Editar el nombre, desactivar y volver a activar. Confirmar que el estado cambie sin duplicar registros.",
        "Abrir Historial y comprobar que los eventos previos siguen visibles; una caída debe decir CAÍDA y verse en rojo.",
        "Mantener Telegram opcional y verificar que sus contactos sigan en Configuración, separados de la agenda telefónica.",
    ])
    doc.add_heading("Carga segura del firmware", level=2)
    add_numbers(doc, [
        "Conectar el ESP32 por USB y cerrar monitores seriales que ocupen el puerto.",
        "Ejecutar npm run firmware:ports para identificar el puerto exacto.",
        "Compilar otra vez con npm run firmware:biosys:build.",
        "Ejecutar npm run firmware:biosys:upload. Pulsar y mantener IO0/BOOT sólo si esptool queda esperando conexión; soltarlo cuando comience Writing.",
        "No desconectar durante la escritura. Abrir el monitor serial y registrar la línea de versión BIOSYS 1.0.7.",
    ])
    doc.add_heading("Prueba de Mensajería", level=2)
    add_numbers(doc, [
        "Después del arranque, comprobar que MENSAJERÍA sea la primera opción y que las otras cuatro sigan presentes.",
        "Entrar en MENSAJE, seleccionar el contacto de prueba y confirmar. La TFT debe mostrar Pendiente/Aceptado, no Enviado.",
        "Comprobar en la app que aparece MENSAJE con el contacto correcto. Tocar el evento y Abrir mensaje; verificar que el compositor se abre sin envío automático.",
        "Repetir con LLAMADA. Abrir marcador y cancelar antes de llamar si no se desea completar la comunicación.",
        "Desactivar el contacto en la app. Esperar sincronización y comprobar que desaparece de la lista TFT.",
    ])
    doc.add_heading("Push Android e iOS", level=2)
    add_numbers(doc, [
        "Con la app en segundo plano, bloquear el teléfono y crear un evento de prueba.",
        "Confirmar que la pantalla sólo muestre VitalWatch y texto genérico; no debe mostrar teléfono, nombre ni signos.",
        "Tocar la notificación: debe abrir el evento exacto. Repetir después de cerrar sesión; iniciar sesión y confirmar que se conserva el destino.",
        "Ejecutar el procedimiento completo una vez en Android y otra vez en iOS; registrar modelo, versión del SO y resultado por separado.",
    ])
    doc.add_heading("Prueba negativa de autorización", level=2)
    add_paragraph(doc, "Antes de promover la candidata, crear dos cuentas de prueba. Un contacto de la cuenta B no debe poder usarse al construir una solicitud autenticada por el dispositivo de A. Documentar HTTP, mensaje y ausencia de evento. Eliminar luego sólo los datos de prueba identificados.")

    part(doc, 13, "Problemas encontrados")
    add_table(doc, ["Área", "Problema", "Impacto", "Resolución o decisión"], [
        ("Git", "Objetos faltantes y árbol con cambios previos", "No permite confiar en commits para rollback", "Backup explícito; no reset ni descarte"),
        ("EAS", "El primer intento intentó clonar el Git dañado", "No llegó a compilar y consumió versionCode 16", "Repetición controlada con EAS_NO_VCS; build 17 finalizado"),
        ("Supabase CLI", "Login de CLI Bun 1.4.0 produjo segmentation fault", "No bloqueó backend", "Token guardado de forma segura y despliegue web para funciones"),
        ("Permiso Edge", "CLI respondió 401 al desplegar funciones", "Impidió deploy por CLI", "Despliegue autorizado desde Dashboard y verificación de timestamp"),
        ("Firmware", "Flash ya usa 89 %", "Margen reducido para futuras funciones", "Módulo compacto; documentar optimización futura"),
        ("Hardware", "ESP32 y teléfonos no estuvieron bajo prueba integral", "No hay validación física", "Mantener candidata, no baseline"),
        ("Offline", "Cola de mensajería vive en RAM", "Reinicio puede perder una solicitud pendiente", "Mostrar estados honestos; evaluar NVS después"),
        ("Expo", "Push Android no se prueba en Expo Go con SDK moderno", "Requiere build nativo", "APK interno 1.0.9 build 17 generado"),
    ], widths=[1.0, 2.25, 1.7, 2.35], font_size=7.8)
    add_paragraph(doc, "Ninguno de estos imprevistos obligó a reescribir la arquitectura ni a tocar sensores. Sí impiden presentar el candidato como producto final probado.")

    part(doc, 14, "Pendientes")
    add_bullets(doc, [
        "Cargar BIOSYS 1.0.7 en el ESP32 y registrar arranque, menú, botones, WiFi y solicitudes.",
        "Instalar App 1.0.9 desde el APK candidato build 17 y repetir pruebas Android reales.",
        "Generar un build iOS firmado y repetir de forma independiente.",
        "Probar lock screen, toque con app cerrada, login intermedio, tel: y sms: en ambos sistemas.",
        "Ensayar aislamiento con dos cuentas y dos contactos reales de prueba.",
        "Ensayar idempotencia enviando dos veces el mismo source_event_id y comprobar un solo device_event.",
        "Evaluar persistencia NVS de la cola sólo si el riesgo de pérdida tras reinicio lo justifica.",
        "Medir tamaño por símbolos y recortar cadenas/recursos antes de nuevas funciones; no optimizar algoritmos médicos sin auditoría separada.",
        "Reparar Git en una tarea controlada después de salvaguardar todos los cambios anteriores del usuario.",
        "Promover a baseline únicamente con evidencia física firmada por el equipo; hasta entonces conservar CANDIDATE FOR REVIEW.",
    ])
    doc.add_heading("Criterio de salida", level=2)
    add_paragraph(doc, "La candidata puede considerarse lista para revisión de laboratorio cuando complete las pruebas físicas sin regresión. Una validación clínica, certificación médica o evaluación regulatoria requiere un proceso diferente y no queda implícita por este trabajo de software.")

    part(doc, 15, "Rollback")
    add_paragraph(doc, "El rollback fue preparado antes del delta porque el repositorio Git no es confiable en su estado actual. No se debe ejecutar git reset --hard ni borrar el árbol de trabajo.")
    add_table(doc, ["Recurso", "Ubicación o hash"], [
        ("Backup anterior", "docs/rollback/VitalWatch_Mensajeria_1_0_9_before.zip"),
        ("SHA-256 backup", "7905AAC3849A7D1CC05B42B50DA1625C12A435A5A3E4DFCDB7DF6BCBC7687115"),
        ("Base firmware inmediata", "esp32/VitalWatch_BIOSYS_1_0_6"),
        ("Baseline física", "esp32/VitalWatch_BIOSYS_1_0_5"),
        ("Paquete candidato sin secreto", "docs/VitalWatch_BIOSYS_1_0_7_Arduino.zip"),
        ("SHA-256 paquete", "62A58348801A9CCB55BD8B358B0C9D1DDA0A1C8C7CF9FAC634B045C979D3A395"),
        ("APK Android candidato", "docs/VitalWatch_APP_1.0.9_CANDIDATE_build17.apk"),
        ("SHA-256 APK", "AA8DA3452FDC4FB5908A7DAD6A680F1FB82ED4F6742EACFE3B59081878BC7E2B"),
    ], widths=[2.05, 5.05], font_size=8.3)
    doc.add_heading("Rollback de firmware", level=2)
    add_numbers(doc, [
        "Si 1.0.7 falla, detener la prueba y conservar logs/fotos.",
        "Compilar y cargar la carpeta BIOSYS 1.0.5, o 1.0.6 sólo si el objetivo es comparar candidatas.",
        "No copiar vitalwatch_config.h a paquetes públicos; usar la configuración privada local controlada.",
        "Confirmar por Serial la versión restaurada y repetir navegación/sensores básicos.",
    ])
    doc.add_heading("Rollback de app", level=2)
    add_numbers(doc, [
        "Desinstalar la candidata 1.0.9 sólo si la prueba lo requiere y reinstalar el último binario aprobado disponible.",
        "Para revertir fuentes, extraer el backup en una carpeta temporal y comparar archivo por archivo; no sobrescribir a ciegas cambios posteriores.",
        "Repetir TypeScript, lint y pruebas antes de construir el binario restaurado.",
    ])
    doc.add_heading("Rollback de Supabase", level=2)
    add_paragraph(doc, "No retirar primero la base mientras existan clientes 1.0.7. El orden seguro es volver firmware/app y después, en una ventana controlada, desactivar la función nueva. La columna contact_id es nullable y compatible con eventos anteriores; dejarla sin uso es más seguro que borrarla de inmediato. Revertir migraciones destructivamente requiere respaldo de base y aprobación explícita.")
    doc.add_heading("Cierre", level=2)
    add_paragraph(doc, "La entrega conserva VitalWatch como un solo sistema coherente: App 1.0.9 candidata, BIOSYS 1.0.7 con SYS 0.9.7 y BIO 0.6.3, y Supabase vigente. La implementación está preparada para revisión y prueba real, no declarada estable.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Registro de implementación de Mensajería en VitalWatch"
    doc.core_properties.subject = "App 1.0.9 y BIOSYS 1.0.7 candidatos"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
